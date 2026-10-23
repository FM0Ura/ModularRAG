import io
from typing import List
import docx
from app.domain.interfaces import DocumentParser
from app.domain.models import RawDocumentPage


class DocxParser(DocumentParser):
    def parse(self, file_bytes: bytes, filename: str) -> List[RawDocumentPage]:
        if not file_bytes or len(file_bytes) == 0:
            raise ValueError(f"Arquivo vazio: '{filename}' não contém dados.")

        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Também extrai conteúdo de tabelas
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        if not full_text.strip():
            raise ValueError(f"Documento DOCX '{filename}' está vazio.")

        return [RawDocumentPage(page_number=1, content=full_text, metadata={"filename": filename})]
