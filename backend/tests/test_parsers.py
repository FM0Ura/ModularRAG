import io
from unittest.mock import MagicMock, patch
import docx
import pytest

from app.adapters.parsers.docx_parser import DocxParser
from app.adapters.parsers.factory import get_parser_for_filename
from app.adapters.parsers.pdf_parser import PDFParser
from app.adapters.parsers.text_parser import TextParser
from app.adapters.parsers.docling_parser import DoclingParser
from app.core.config import Settings


# --- TextParser Tests ---

def test_text_parser_success():
    parser = TextParser()
    sample_text = "Primeira linha.\n\nSegunda linha com informações técnicas."
    pages = parser.parse(sample_text.encode("utf-8"), "documento.txt")
    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert "Segunda linha" in pages[0].content
    assert pages[0].metadata == {"filename": "documento.txt"}


def test_text_parser_empty_file_raises_value_error():
    parser = TextParser()
    with pytest.raises(ValueError, match="Arquivo vazio"):
        parser.parse(b"", "vazio.txt")


def test_text_parser_whitespace_only_raises_value_error():
    parser = TextParser()
    with pytest.raises(ValueError, match="Arquivo vazio"):
        parser.parse(b"   \n  \t  \n", "espacos.txt")


def test_text_parser_latin1_fallback():
    parser = TextParser()
    sample_text = "Atenção: texto codificado em ISO-8859-1 (Latin-1) com caracteres especiais: ç, á, é."
    latin1_bytes = sample_text.encode("latin-1")
    pages = parser.parse(latin1_bytes, "latin1.txt")
    assert len(pages) == 1
    assert "Atenção" in pages[0].content


# --- PDFParser Tests ---

def test_pdf_parser_empty_file_raises_value_error():
    parser = PDFParser()
    with pytest.raises(ValueError, match="Arquivo vazio"):
        parser.parse(b"", "vazio.pdf")


def test_pdf_parser_success_with_page_numbers():
    parser = PDFParser()
    mock_page_1 = MagicMock()
    mock_page_1.extract_text.return_value = "Conteúdo da primeira página."
    mock_page_2 = MagicMock()
    mock_page_2.extract_text.return_value = "   "  # Blank page, should be skipped
    mock_page_3 = MagicMock()
    mock_page_3.extract_text.return_value = "Conteúdo da terceira página."

    with patch("pypdf.PdfReader") as mock_reader_cls:
        mock_reader = MagicMock()
        mock_reader.pages = [mock_page_1, mock_page_2, mock_page_3]
        mock_reader_cls.return_value = mock_reader

        pages = parser.parse(b"%PDF-1.4 mock content", "relatorio.pdf")

    assert len(pages) == 2
    assert pages[0].page_number == 1
    assert pages[0].content == "Conteúdo da primeira página."
    assert pages[0].metadata == {"filename": "relatorio.pdf"}

    assert pages[1].page_number == 3
    assert pages[1].content == "Conteúdo da terceira página."
    assert pages[1].metadata == {"filename": "relatorio.pdf"}


def test_pdf_parser_no_readable_text_raises_value_error():
    parser = PDFParser()
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "   \n "

    with patch("pypdf.PdfReader") as mock_reader_cls:
        mock_reader = MagicMock()
        mock_reader.pages = [mock_page]
        mock_reader_cls.return_value = mock_reader

        with pytest.raises(ValueError, match="Não foi possível extrair texto legível"):
            parser.parse(b"%PDF-1.4 empty text", "sem_texto.pdf")


# --- DocxParser Tests ---

def test_docx_parser_empty_file_raises_value_error():
    parser = DocxParser()
    with pytest.raises(ValueError, match="Arquivo vazio"):
        parser.parse(b"", "vazio.docx")


def test_docx_parser_success_with_paragraphs_and_tables():
    parser = DocxParser()
    doc = docx.Document()
    doc.add_paragraph("Introdução ao documento.")
    doc.add_paragraph("Segundo parágrafo com detalhes.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"
    table.rows[1].cells[0].text = "Valor A"
    table.rows[1].cells[1].text = "Valor B"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    pages = parser.parse(docx_bytes, "documento.docx")
    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert "Introdução ao documento." in pages[0].content
    assert "Segundo parágrafo com detalhes." in pages[0].content
    assert "Header 1 | Header 2" in pages[0].content
    assert "Valor A | Valor B" in pages[0].content
    assert pages[0].metadata == {"filename": "documento.docx"}


def test_docx_parser_empty_document_raises_value_error():
    parser = DocxParser()
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    with pytest.raises(ValueError, match="está vazio"):
        parser.parse(docx_bytes, "em_branco.docx")


# --- Factory Tests ---

def test_factory_selects_correct_parser():
    p1 = get_parser_for_filename("doc.pdf")
    p2 = get_parser_for_filename("doc.docx")
    p3 = get_parser_for_filename("doc.md")
    p4 = get_parser_for_filename("doc.txt")
    p5 = get_parser_for_filename("DOC.PDF")
    # Default is DoclingParser for rich PDF/DOCX layouts
    assert p1.__class__.__name__ == "DoclingParser"
    assert p2.__class__.__name__ == "DoclingParser"
    assert p3.__class__.__name__ == "TextParser"
    assert p4.__class__.__name__ == "TextParser"
    assert p5.__class__.__name__ == "DoclingParser"


def test_factory_unsupported_extension_raises_value_error():
    with pytest.raises(ValueError, match="Formato de arquivo não suportado"):
        get_parser_for_filename("planilha.xlsx")


# --- DoclingParser Tests ---

def test_docling_parser_empty_file_raises_value_error():
    parser = DoclingParser()
    with pytest.raises(ValueError, match="Arquivo vazio"):
        parser.parse(b"", "vazio.pdf")


def test_docling_parser_mocked_success_with_pages():
    mock_converter = MagicMock()
    mock_result = MagicMock()
    mock_doc = MagicMock()
    mock_doc.pages = {1: MagicMock(), 2: MagicMock()}
    mock_doc.export_to_markdown.side_effect = lambda page_no=None: f"Conteudo da pagina {page_no}" if page_no else "Conteudo total"
    mock_result.document = mock_doc
    mock_converter.convert.return_value = mock_result

    parser = DoclingParser(converter=mock_converter)
    pages = parser.parse(b"simulacao de bytes", "documento.pdf")

    assert len(pages) == 2
    assert pages[0].page_number == 1
    assert "Conteudo da pagina 1" in pages[0].content
    assert pages[1].page_number == 2
    assert "Conteudo da pagina 2" in pages[1].content
    assert pages[0].metadata == {"filename": "documento.pdf"}


def test_docling_parser_real_markdown_conversion():
    parser = DoclingParser()
    sample_content = b"# Documento Tecnico\n\nTexto de teste com tabela:\n\n| Parametro | Valor |\n|---|---|\n| Timeout | 45s |"
    pages = parser.parse(sample_content, "tabela.txt")

    assert len(pages) >= 1
    assert pages[0].page_number == 1
    assert "Documento Tecnico" in pages[0].content
    assert "Timeout" in pages[0].content
    assert "45s" in pages[0].content


def test_factory_selects_docling_when_configured():
    custom_settings = Settings(PARSER_PROVIDER="docling")
    p_pdf = get_parser_for_filename("doc.pdf", config=custom_settings)
    p_docx = get_parser_for_filename("doc.docx", config=custom_settings)
    assert p_pdf.__class__.__name__ == "DoclingParser"
    assert p_docx.__class__.__name__ == "DoclingParser"

    std_settings = Settings(PARSER_PROVIDER="standard")
    p_std_pdf = get_parser_for_filename("doc.pdf", config=std_settings)
    p_std_docx = get_parser_for_filename("doc.docx", config=std_settings)
    assert p_std_pdf.__class__.__name__ == "PDFParser"
    assert p_std_docx.__class__.__name__ == "DocxParser"

